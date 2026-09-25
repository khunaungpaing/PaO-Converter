"""DOCX-to-DOCX Pa-O conversion with source-font filtering."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Callable, Iterator

from core.cancellation import raise_if_cancelled
from core.engine import convert_pao_ascii_to_unicode
from core.file_options import DEFAULT_SIZE_MAPPING, font_names_match, mapped_font_size


def _iter_table_paragraphs(table: Any) -> Iterator[Any]:
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested_table in cell.tables:
                yield from _iter_table_paragraphs(nested_table)


def _iter_docx_paragraphs(doc: Any) -> Iterator[Any]:
    """Yield editable paragraphs from the body, tables, headers, and footers."""
    seen: set[int] = set()

    def unique(paragraphs: Iterator[Any]) -> Iterator[Any]:
        for paragraph in paragraphs:
            marker = id(paragraph._p)
            if marker not in seen:
                seen.add(marker)
                yield paragraph

    yield from unique(iter(doc.paragraphs))
    for table in doc.tables:
        yield from unique(_iter_table_paragraphs(table))

    story_names = (
        "header", "first_page_header", "even_page_header",
        "footer", "first_page_footer", "even_page_footer",
    )
    for section in doc.sections:
        for story_name in story_names:
            story = getattr(section, story_name)
            if story.is_linked_to_previous:
                continue
            yield from unique(iter(story.paragraphs))
            for table in story.tables:
                yield from unique(_iter_table_paragraphs(table))


def _normal_style_font(doc: Any) -> Any | None:
    try:
        return doc.styles["Normal"].font
    except (KeyError, TypeError):
        return None


def _run_font_name(run: Any, paragraph: Any, doc: Any) -> str | None:
    """Resolve the most relevant direct or inherited DOCX font family."""
    candidates = [
        run.font,
        getattr(getattr(run, "style", None), "font", None),
        getattr(getattr(paragraph, "style", None), "font", None),
        _normal_style_font(doc),
    ]
    for font in candidates:
        name = getattr(font, "name", None)
        if name:
            return name
    return None


def _run_font_size(run: Any, paragraph: Any, doc: Any) -> float | None:
    """Resolve a run's direct or inherited point size."""
    fonts = [
        run.font,
        getattr(getattr(run, "style", None), "font", None),
        getattr(getattr(paragraph, "style", None), "font", None),
        _normal_style_font(doc),
    ]
    for font in fonts:
        size = getattr(font, "size", None)
        if size is not None:
            return float(size.pt)
    return None


def convert_docx_file(
    src_path: str | Path,
    dst_path: str | Path,
    source_font: str | None,
    output_font_family: str,
    size_mapping: dict[float, float] | None = None,
    progress: Callable[[int], None] | None = None,
    cancelled: Callable[[], bool] | None = None,
) -> int:
    """Convert adjacent matching runs together, using the first run's format.

    Return the number of source runs converted.
    """
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    src = Path(src_path)
    dst = Path(dst_path)
    if not src.is_file():
        raise FileNotFoundError(f"Source DOCX not found: {src}")

    sizes = DEFAULT_SIZE_MAPPING if size_mapping is None else size_mapping
    doc = Document(str(src))
    paragraphs = list(_iter_docx_paragraphs(doc))
    total = len(paragraphs) or 1
    converted_runs = 0

    for index, paragraph in enumerate(paragraphs, 1):
        raise_if_cancelled(cancelled)
        groups: list[list[Any]] = []
        current: list[Any] = []
        for run in paragraph.runs:
            actual_font = _run_font_name(run, paragraph, doc)
            if font_names_match(actual_font, source_font):
                if run.text:
                    current.append(run)
            else:
                if current:
                    groups.append(current)
                    current = []
        if current:
            groups.append(current)

        for group in groups:
            run = group[0]
            original_size = _run_font_size(run, paragraph, doc)
            text = "".join(part.text for part in group)
            run.text = convert_pao_ascii_to_unicode(text)
            for part in group[1:]:
                part.text = ""
            run.font.name = output_font_family
            r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
            for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                r_fonts.set(qn(f"w:{attribute}"), output_font_family)
            if original_size is not None:
                run.font.size = Pt(mapped_font_size(original_size, sizes))
            converted_runs += len(group)

        if progress:
            progress(int(index / total * 90))

    raise_if_cancelled(cancelled)
    doc.save(str(dst))
    if progress:
        progress(100)
    return converted_runs
