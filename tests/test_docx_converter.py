"""Integration tests for selective DOCX conversion."""

import os
import tempfile
import unittest

try:
    from docx import Document
    from docx.shared import Pt
except ImportError:
    Document = None

from migration.docx_converter import convert_docx_file


@unittest.skipIf(Document is None, "python-docx is not installed")
class TestDocxConverter(unittest.TestCase):
    def test_converts_split_syllable_together(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            src = os.path.join(tmp, "source.docx")
            dst = os.path.join(tmp, "converted.docx")
            doc = Document()
            paragraph = doc.add_paragraph()
            first = paragraph.add_run("j")
            first.font.name = "kothupaoh1"
            first.font.size = Pt(16)
            first.bold = True
            second = paragraph.add_run("zGD;")
            second.font.name = "kothupaoh1"
            second.font.size = Pt(20)
            other = paragraph.add_run("keep")
            other.font.name = "Arial"
            last = paragraph.add_run("t")
            last.font.name = "kothupaoh1"
            doc.save(src)

            count = convert_docx_file(src, dst, "kothupaoh1", "KhamThaton")
            runs = Document(dst).paragraphs[0].runs
            self.assertEqual(count, 3)
            self.assertEqual(runs[0].text, "ဖြွီး")
            self.assertEqual(runs[1].text, "")
            self.assertTrue(runs[0].bold)
            self.assertEqual(runs[0].font.size.pt, 12)
            self.assertEqual(runs[2].text, "keep")
            self.assertEqual(runs[2].font.name, "Arial")
            self.assertEqual(runs[3].text, "အ")

    def test_converts_only_selected_font_and_maps_size(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            src = os.path.join(tmp, "source.docx")
            dst = os.path.join(tmp, "converted.docx")
            doc = Document()
            paragraph = doc.add_paragraph()
            selected = paragraph.add_run("t")
            selected.font.name = "kothupaoh1"
            selected.font.size = Pt(16)
            untouched = paragraph.add_run(" keep")
            untouched.font.name = "Arial"
            untouched.font.size = Pt(16)
            doc.save(src)

            count = convert_docx_file(
                src,
                dst,
                source_font="Kothu PaOh1",
                output_font_family="KhamThaton",
                size_mapping={16.0: 12.0},
            )

            self.assertEqual(count, 1)
            result = Document(dst)
            runs = result.paragraphs[0].runs
            self.assertEqual(runs[0].text, "အ")
            self.assertEqual(runs[0].font.name, "KhamThaton")
            self.assertAlmostEqual(runs[0].font.size.pt, 12.0)
            self.assertEqual(runs[1].text, " keep")
            self.assertEqual(runs[1].font.name, "Arial")
